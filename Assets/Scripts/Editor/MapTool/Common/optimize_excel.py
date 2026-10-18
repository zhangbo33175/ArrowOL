import openpyxl
import os

# 配置表路径
base_path = r"E:\Game\CatGame\CatGameCodeOptimizationBuild\Docs\Designs\Excels\Tables\AlTables\Item"

# 列名映射：中文列名 -> 英文字段名
column_mapping = {
    "ID": "TemplateId",
    "模板名称": "TemplateName",
    "建筑名称": "TemplateName",
    "物件名称": "TemplateName",
    "ResName": "ResName",
    "Icon": "ImagePath",
    "默认动画": "AnimName",
    "动画名": "AnimName",
    "是否循环": "Loop",
    "播放速度": "Speed",
    "是否可交互": "Interactable",
    "交互类型": "InteractType",
    "渲染层级": "SortOrder",
    "是否有碰撞": "CollisionEnabled",
    "对话气泡": "DialogBubble",
    "初始朝向": "FacingDirection",
    "朝向": "FacingDirection",
    "是否显示阴影": "ShadowEnabled",
    "是否收集物": "IsCollectible",
    "重生时间": "RespawnTime",
}

files = ["TableBuildings.xlsm", "TableCharacters.xlsm", "TableProps.xlsm", "TableTemplate.xlsm"]

for filename in files:
    filepath = os.path.join(base_path, filename)
    if not os.path.exists(filepath):
        print(f"跳过不存在的文件: {filename}")
        continue

    print(f"\n处理: {filename}")

    # 打开Excel
    wb = openpyxl.load_workbook(filepath)
    ws = wb.active

    # 打印当前列名（第2行）
    print("  原始列名:")
    for col in range(1, ws.max_column + 1):
        col_name = ws.cell(row=2, column=col).value
        print(f"    列{col}: {col_name}")

    # 修改列名（第2行）
    print("  修改后列名:")
    for col in range(1, ws.max_column + 1):
        old_name = ws.cell(row=2, column=col).value
        if old_name in column_mapping:
            new_name = column_mapping[old_name]
            ws.cell(row=2, column=col).value = new_name
            print(f"    列{col}: {old_name} -> {new_name}")
        else:
            print(f"    列{col}: {old_name} (未匹配，保留)")

    # 保存
    wb.save(filepath)
    print(f"  已保存: {filename}")

print("\n全部完成！")
